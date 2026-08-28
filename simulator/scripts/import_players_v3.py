#!/usr/bin/env python3
"""players-v3-4attrs importer: All players + Card (eligible positions, skill moves)
+ Contracts (club/wage/value) sheets → data/players.json.

Conventions preserved from v2: sequential player ids in workbook rank order;
attribute_stats over the full population for outfield attributes and over the
GK population for gk_* attributes (the locked GK-normalization fix). The four
new attributes (+ curve, imported as data with no engine role this phase) use
the full population, consistent with every other outfield attribute.
"""
from __future__ import annotations

import argparse
import json
import statistics as st
from pathlib import Path

import openpyxl

ATTRIBUTE_COLUMNS = {
    "acceleration": "Acceleration", "sprint_speed": "Sprint Speed", "agility": "Agility",
    "reactions": "Reactions", "ball_control": "Ball Control", "dribbling": "Dribbling",
    "short_passing": "Short Passing", "long_passing": "Long Passing", "vision": "Vision",
    "crossing": "Crossing", "finishing": "Finishing", "attacking_position": "Att. Position",
    "shot_power": "Shot Power", "long_shots": "Long Shots", "volleys": "Volleys",
    "heading_accuracy": "Heading Acc.", "defensive_awareness": "Def. Awareness",
    "standing_tackle": "Standing Tackle", "sliding_tackle": "Sliding Tackle",
    "interceptions": "Interceptions", "strength": "Strength", "stamina": "Stamina",
    "aggression": "Aggression", "jumping": "Jumping",
    "gk_diving": "GK Diving", "gk_handling": "GK Handling", "gk_kicking": "GK Kicking",
    "gk_positioning": "GK Positioning", "gk_reflexes": "GK Reflexes",
    # players-v3 additions
    "free_kick_accuracy": "FK Accuracy", "penalties": "Penalties",
    "composure": "Composure", "balance": "Balance", "curve": "Curve",
}
GK_POPULATION_ATTRS = {"gk_diving", "gk_handling", "gk_kicking", "gk_positioning", "gk_reflexes"}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--out", type=Path, default=Path("data/players.json"))
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb["All players"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h}

    card = {}
    cws = wb["Card"]
    cheaders = [c.value for c in cws[1]]
    cidx = {h: i for i, h in enumerate(cheaders) if h}
    for r in cws.iter_rows(min_row=2, values_only=True):
        if r[cidx["Player"]]:
            card[str(r[cidx["Player"]])] = {
                "eligible": str(r[cidx["Eligible positions"]] or ""),
                "skill_moves": int(r[cidx["SM"]] or 3),
                "face": [r[cidx[c]] for c in ("PAC / DIV", "SHO / HAN", "PAS / KIC", "DRI / REF", "DEF / SPD", "PHY / POS")],
            }

    contracts = {}
    tws = wb["Contracts"]
    theaders = [c.value for c in tws[1]]
    tidx = {h: i for i, h in enumerate(theaders) if h}
    for r in tws.iter_rows(min_row=2, values_only=True):
        if r[tidx["Player"]]:
            contracts[str(r[tidx["Player"]])] = {
                "club": r[tidx["Club"]], "national_team": r[tidx["National team"]],
                "contract_until": r[tidx["Contract until"]], "years_left": r[tidx["Years left"]],
                "wage_eur_wk": r[tidx["Wage (EUR/wk)"]], "value_eur_m": r[tidx["Value (EUR m)"]],
            }

    players = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = r[idx["Player"]]
        pos = r[idx["Position"]]
        if not name or not pos:
            continue
        name = str(name)
        attrs = {}
        for key, col in ATTRIBUTE_COLUMNS.items():
            v = r[idx[col]]
            if isinstance(v, (int, float)):
                attrs[key] = float(v)
        c = card.get(name, {})
        elig = [p.strip() for p in c.get("eligible", "").split("/") if p.strip()]
        secondary = [p for p in elig if p != str(pos).upper()]
        k = contracts.get(name, {})
        players.append({
            "player_id": f"p{len(players)+1:03d}",
            "name": name,
            "primary_position": str(pos).upper(),
            "secondary_positions": secondary,
            "preferred_foot": str(r[idx["Foot"]] or "R").upper(),
            "weak_foot": int(r[idx["WF"]] or 3),
            "skill_moves": c.get("skill_moves", 3),
            "natural_side": str(r[idx["Nat. side"]] or "C").upper(),
            "height_cm": float(r[idx["Ht (cm)"]]) if isinstance(r[idx["Ht (cm)"]], (int, float)) else None,
            "weight_kg": float(r[idx["Wt (kg)"]]) if isinstance(r[idx["Wt (kg)"]], (int, float)) else None,
            # Presentation/selection metadata ONLY — never used in event resolution.
            "ovr": int(r[idx["OVR"]] or 0),
            "pot": int(r[idx["POT"]] or 0),
            "age": int(r[idx["Age"]] or 0),
            "club": k.get("club"), "national_team": k.get("national_team"),
            "contract_until": k.get("contract_until"), "wage_eur_wk": k.get("wage_eur_wk"),
            "value_eur_m": k.get("value_eur_m"),
            "card_face": c.get("face"),
            "attributes": attrs,
            "source_row": len(players) + 2,
        })

    gk_players = [p for p in players if p["primary_position"] == "GK"]
    stats = {}
    for key in ATTRIBUTE_COLUMNS:
        pop = gk_players if key in GK_POPULATION_ATTRS else players
        vals = [p["attributes"][key] for p in pop if key in p["attributes"]]
        if not vals:
            continue
        entry = {"mean": st.mean(vals), "sd": st.stdev(vals) if len(vals) > 1 else 1.0, "n": len(vals)}
        if key in GK_POPULATION_ATTRS:
            entry["population"] = "GK"
        stats[key] = entry

    payload = {
        "source_workbook": args.workbook.name,
        "source_sheet": "All players",
        "columns": [h for h in headers if h],
        "height_detected": True, "weight_detected": True,
        "players": players,
        "attribute_stats": stats,
        "data_version": "players-v3-4attrs",
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(json.dumps({"players": len(players), "gk": len(gk_players),
                      "attributes": len(ATTRIBUTE_COLUMNS), "out": str(args.out)}))


if __name__ == "__main__":
    main()
